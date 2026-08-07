from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, status
from motor.motor_asyncio import AsyncIOMotorDatabase
import openpyxl
import io
from typing import Dict, List
from datetime import datetime, timezone

from app.core.database import get_db
from app.modules.auth.service import require_role, create_user
from app.modules.auth.schemas import UserCreate
from app.modules.library.schemas import BookCreate
from app.modules.library.router import create_book

router = APIRouter(prefix="/bulk-upload", tags=["Bulk Uploads"])

STUDENT_HEADERS = ["Name", "Email", "Password", "Role", "BatchName"]
BOOK_HEADERS = ["Title", "Author", "ISBN", "Quantity"]

@router.get("/schemas", response_model=Dict[str, List[str]])
async def get_upload_schemas(
    current_user = Depends(require_role(["head"]))
):
    """Area Head retrieves expected schemas and columns prior to excel upload."""
    return {
        "students": STUDENT_HEADERS,
        "books": BOOK_HEADERS
    }

@router.post("/students")
async def bulk_upload_students(
    file: UploadFile = File(...),
    db: AsyncIOMotorDatabase = Depends(get_db),
    current_user = Depends(require_role(["head"]))
):
    """
    Bulk uploads students from an XLSX file.
    Validates columns layout prior to inserting.
    """
    try:
        contents = await file.read()
        workbook = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
        sheet = workbook.active
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to parse Excel file: {str(e)}")

    # Header check
    header_row = [cell.value for cell in sheet[1]]
    for h in STUDENT_HEADERS:
        if h not in header_row:
            raise HTTPException(
                status_code=400,
                detail=f"Columns mismatch. Expected: {STUDENT_HEADERS}. Found: {header_row}"
            )

    header_indices = {name: header_row.index(name) for name in STUDENT_HEADERS}
    imported = 0
    errors = []

    for r_idx in range(2, sheet.max_row + 1):
        row = sheet[r_idx]
        if not any(cell.value is not None for cell in row):
            continue
            
        try:
            name = str(row[header_indices["Name"]].value or "").strip()
            email = str(row[header_indices["Email"]].value or "").strip().lower()
            password = str(row[header_indices["Password"]].value or "").strip()
            role = str(row[header_indices["Role"]].value or "student").strip().lower()
            batch_name = str(row[header_indices["BatchName"]].value or "").strip()

            if not name or not email or not password:
                errors.append(f"Row {r_idx}: Name, Email, and Password are required.")
                continue

            # Ensure batch exists in DB, or create it if not
            if batch_name:
                batch_exists = await db.batches.find_one({"name": batch_name})
                if not batch_exists:
                    await db.batches.insert_one({
                        "name": batch_name,
                        "description": f"Automatically generated during bulk import",
                        "is_open": False,
                        "created_at": datetime.now(timezone.utc)
                    })

            user_in = UserCreate(
                email=email,
                name=name,
                password=password,
                role=role,
                batch_id=batch_name
            )
            
            await create_user(db, user_in)
            imported += 1
        except Exception as e:
            errors.append(f"Row {r_idx}: Error - {str(e)}")

    return {
        "message": f"Successfully imported {imported} users.",
        "errors": errors
    }

@router.post("/books")
async def bulk_upload_books(
    file: UploadFile = File(...),
    db: AsyncIOMotorDatabase = Depends(get_db),
    current_user = Depends(require_role(["head"]))
):
    """
    Bulk uploads library books from an XLSX file.
    Validates columns layout prior to inserting.
    Generates copy IDs sequentially.
    """
    try:
        contents = await file.read()
        workbook = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
        sheet = workbook.active
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to parse Excel file: {str(e)}")

    # Header check
    header_row = [cell.value for cell in sheet[1]]
    for h in BOOK_HEADERS:
        if h not in header_row:
            raise HTTPException(
                status_code=400,
                detail=f"Columns mismatch. Expected: {BOOK_HEADERS}. Found: {header_row}"
            )

    header_indices = {name: header_row.index(name) for name in BOOK_HEADERS}
    imported = 0
    errors = []

    for r_idx in range(2, sheet.max_row + 1):
        row = sheet[r_idx]
        if not any(cell.value is not None for cell in row):
            continue
            
        try:
            title = str(row[header_indices["Title"]].value or "").strip()
            author = str(row[header_indices["Author"]].value or "").strip()
            isbn = str(row[header_indices["ISBN"]].value or "").strip()
            quantity_val = row[header_indices["Quantity"]].value
            
            try:
                quantity = int(quantity_val)
            except (ValueError, TypeError):
                errors.append(f"Row {r_idx}: Invalid quantity '{quantity_val}'. Must be an integer.")
                continue

            if not title or not author or not isbn or quantity <= 0:
                errors.append(f"Row {r_idx}: Title, Author, ISBN are required, and Quantity must be > 0.")
                continue

            # Verify ISBN unique in DB
            existing = await db.books.find_one({"isbn": isbn})
            if existing:
                errors.append(f"Row {r_idx}: Book with ISBN '{isbn}' already exists.")
                continue

            # Construct book
            copies = []
            for i in range(1, quantity + 1):
                copies.append({
                    "copy_id": f"{isbn}-{i}",
                    "status": "available",
                    "lent_to": None
                })
                
            book_dict = {
                "title": title,
                "author": author,
                "isbn": isbn,
                "quantity": quantity,
                "copies": copies
            }
            
            await db.books.insert_one(book_dict)
            imported += 1
        except Exception as e:
            errors.append(f"Row {r_idx}: Error - {str(e)}")

    return {
        "message": f"Successfully imported {imported} library books.",
        "errors": errors
    }
