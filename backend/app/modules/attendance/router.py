from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form, status
from motor.motor_asyncio import AsyncIOMotorDatabase
from typing import List, Dict, Any, Optional
import openpyxl
import io
from datetime import datetime, timezone

from app.core.database import get_db
from app.modules.auth.service import require_role, get_current_user
from app.modules.attendance.schemas import (
    AttendanceRecord, 
    StudentAttendanceSummary, 
    AttendanceManualUpdateRequest
)

def safe_int(val, default=0):
    if val is None or str(val).strip() == "":
        return default
    try:
        return int(float(str(val).strip()))
    except (ValueError, TypeError):
        return default

def safe_float(val, default=0.0):
    if val is None or str(val).strip() == "":
        return default
    try:
        return float(str(val).strip())
    except (ValueError, TypeError):
        return default

def parse_date_to_str(d_val) -> Optional[str]:
    if not d_val:
        return None
    if isinstance(d_val, datetime):
        return d_val.strftime("%Y-%m-%d")
    
    # Try parsing string
    val_str = str(d_val).strip()
    if not val_str:
        return None
        
    # Common formats to try (preferring DD/MM/YYYY over MM/DD/YYYY)
    formats = [
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d",
        "%d-%m-%Y",
        "%d/%m/%Y",
        "%d-%b-%Y",
        "%d-%b-%y",
        "%d-%B-%Y",
        "%Y/%m/%d",
        "%m/%d/%Y",
    ]
    for fmt in formats:
        try:
            dt = datetime.strptime(val_str, fmt)
            return dt.strftime("%Y-%m-%d")
        except ValueError:
            continue
            
    return val_str


router = APIRouter(prefix="/attendance", tags=["Attendance Management"])

EXPECTED_HEADERS_OLD = ["StudentEmail", "Date", "Session 1", "Session 2", "Session 3", "Session 4"]
EXPECTED_HEADERS_NEW = ["Sr No.", "Asp ID", "Name", "Skillfy App Downloaded or Not", "Total Sessions", "Sessions Attended", "S-1 to S-4 (Repeating)"]

@router.get("/schema", response_model=Dict[str, List[str]])
async def get_attendance_schema(current_user = Depends(get_current_user)):
    """Returns the expected column headers and order for XLSX attendance sheets."""
    return {"expected_columns": EXPECTED_HEADERS_NEW}

@router.post("/upload")
async def upload_attendance(
    batch_id: str = Form(...),
    file: UploadFile = File(...),
    db: AsyncIOMotorDatabase = Depends(get_db),
    current_user = Depends(require_role(["head", "trainer"]))
):
    """
    Upload an XLSX spreadsheet to update attendance.
    Supports both the old format (StudentEmail, Date, etc.) and the new multi-column format.
    Trainers can only upload for batches assigned to them.
    """
    # 1. Authorize trainer check
    user_role = current_user.get("role")
    today_str = datetime.now().strftime("%Y-%m-%d")
    if user_role != "head":
        assigned_classes = current_user.get("classes_assigned") or []
        if batch_id not in assigned_classes:
            raise HTTPException(
                status_code=status.HTTP_403_FORBIDDEN,
                detail=f"You are not assigned to class/batch '{batch_id}' and cannot upload attendance for it."
            )
            
    # Verify batch exists
    batch = await db.batches.find_one({"name": batch_id})
    if not batch:
        # Fallback or log, but we continue
        pass

    # 2. Parse XLSX
    try:
        contents = await file.read()
        workbook = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
        sheet = workbook.active
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to parse Excel file: {str(e)}")

    # 3. Detect format
    first_row = [str(cell.value).strip() if cell.value is not None else "" for cell in sheet[1]]
    is_old_format = any(h == "StudentEmail" for h in first_row)

    if is_old_format:
        # --- OLD FORMAT PARSING ---
        for expected in EXPECTED_HEADERS_OLD:
            if expected not in first_row:
                raise HTTPException(
                    status_code=400,
                    detail=f"Columns mismatch. Expected: {EXPECTED_HEADERS_OLD}. Found: {first_row}"
                )

        header_indices = {name: first_row.index(name) for name in EXPECTED_HEADERS_OLD}
        records_added = 0
        errors = []

        for r_idx in range(2, sheet.max_row + 1):
            row = sheet[r_idx]
            if not any(cell.value is not None for cell in row):
                continue
                
            try:
                student_email = str(row[header_indices["StudentEmail"]].value or "").strip().lower()
                if not student_email:
                    continue
                date_val = row[header_indices["Date"]].value
                date_str = parse_date_to_str(date_val)
                if not date_str:
                    errors.append(f"Row {r_idx}: Missing or invalid date.")
                    continue

                if user_role != "head" and date_str != today_str:
                    errors.append(f"Row {r_idx}: Skipped date '{date_str}'. Teachers can only update the current date ({today_str}).")
                    continue

                def map_val(v):
                    if v is None or str(v).strip() in ["", "-", "None", "none"]:
                        return "None"
                    v_str = str(v).strip()
                    v_lower = v_str.lower()
                    if v_lower in ["ab", "absent"]:
                        return "Absent"
                    if v_lower == "late":
                        return "Late"
                    if v_lower == "none":
                        return "None"
                    
                    # If it has any alphabets, it is present (return the value itself)
                    if any(c.isalpha() for c in v_str):
                        return v_str
                    return "None"

                session_1 = map_val(row[header_indices["Session 1"]].value)
                session_2 = map_val(row[header_indices["Session 2"]].value)
                session_3 = map_val(row[header_indices["Session 3"]].value)
                session_4 = map_val(row[header_indices["Session 4"]].value)

                student = await db.users.find_one({"email": student_email, "role": "student"})
                if not student:
                    errors.append(f"Row {r_idx}: Student email '{student_email}' does not exist in DB.")
                    continue

                existing = await db.attendance.find_one({
                    "student_email": student_email,
                    "batch_id": batch_id,
                    "date": date_str
                })
                if existing:
                    session_1 = existing.get("session_1", "None") if session_1 == "None" else session_1
                    session_2 = existing.get("session_2", "None") if session_2 == "None" else session_2
                    session_3 = existing.get("session_3", "None") if session_3 == "None" else session_3
                    session_4 = existing.get("session_4", "None") if session_4 == "None" else session_4

                await db.attendance.update_one(
                    {
                        "student_email": student_email,
                        "batch_id": batch_id,
                        "date": date_str
                    },
                    {
                        "$set": {
                            "session_1": session_1,
                            "session_2": session_2,
                            "session_3": session_3,
                            "session_4": session_4,
                            "uploaded_by": current_user["email"]
                        }
                    },
                    upsert=True
                )
                records_added += 1

            except Exception as err:
                errors.append(f"Row {r_idx}: Formatting error. Details: {str(err)}")

        return {
            "message": f"Successfully processed {records_added} attendance records (old format).",
            "warnings_errors": errors
        }

    else:
        # --- NEW FORMAT PARSING ---
        if sheet.max_row < 7:
            raise HTTPException(
                status_code=400,
                detail="Invalid Excel format. Sheet must contain at least 7 rows."
            )
            
        header_row = [str(cell.value).strip() if cell.value is not None else "" for cell in sheet[7]]
        if "Asp ID" not in header_row or "Name" not in header_row:
            raise HTTPException(
                status_code=400,
                detail="Columns mismatch. Expected headers 'Asp ID' and 'Name' in row 7 of the spreadsheet."
            )

        row_5 = [cell.value for cell in sheet[5]]
        
        asp_id_idx = header_row.index("Asp ID")
        name_idx = header_row.index("Name")

        records_added = 0
        summaries_added = 0
        errors = []

        # Dynamically locate columns from Row 4 and Row 7
        skillfy_downloaded_idx = None
        skillfy_courses_completed_idx = None
        total_sessions_idx = None
        sessions_attended_idx = None
        no_of_days_attended_idx = None
        day_wise_attendance_pct_idx = None
        session_wise_attendance_pct_idx = None

        row_4_vals = [str(cell.value).strip() if cell.value is not None else "" for cell in sheet[4]]
        row_7_vals = [str(cell.value).strip() if cell.value is not None else "" for cell in sheet[7]]

        for col_idx in range(len(row_7_vals)):
            r4_val = row_4_vals[col_idx].lower() if col_idx < len(row_4_vals) else ""
            r7_val = row_7_vals[col_idx].lower() if col_idx < len(row_7_vals) else ""

            # Match Skillfy Downloaded (Z or BU)
            if "skillfy app downloaded" in r4_val or "skillfy app downloaded" in r7_val or "aspirant skillfy app" in r4_val:
                skillfy_downloaded_idx = col_idx
            
            # Match Courses Completed (AA or EX)
            if "courses completed" in r4_val or "courses completed" in r7_val or "number of courses completed" in r4_val:
                skillfy_courses_completed_idx = col_idx

            # Match Total Sessions
            if "total sessions" in r4_val or "total sessions" in r7_val:
                total_sessions_idx = col_idx

            # Match Sessions Attended
            if "sessions attended" in r4_val or "sessions attended" in r7_val:
                sessions_attended_idx = col_idx

            # Match No of Days Attended
            if "no of days attended" in r4_val or "no of days attended" in r7_val or "no. of days attended" in r4_val or "no. of days attended" in r7_val:
                no_of_days_attended_idx = col_idx

            # Match Day Wise Attendance %
            if "day wise attendance" in r4_val or "day wise attendance" in r7_val:
                day_wise_attendance_pct_idx = col_idx

            # Match Session Wise Attendance % (including typo "arrendance")
            if "session wise attendance" in r4_val or "session wise attendance" in r7_val or "session wise arrendance" in r4_val or "session wise arrendance" in r7_val:
                session_wise_attendance_pct_idx = col_idx

        # Fallbacks to defaults if not found
        if skillfy_downloaded_idx is None:
            if "Skillfy App Downloaded or Not" in row_7_vals:
                skillfy_downloaded_idx = row_7_vals.index("Skillfy App Downloaded or Not")
            else:
                skillfy_downloaded_idx = 25

        if skillfy_courses_completed_idx is None:
            skillfy_courses_completed_idx = skillfy_downloaded_idx + 1 if skillfy_downloaded_idx is not None else 26
            
        if total_sessions_idx is None:
            total_sessions_idx = 27
        if sessions_attended_idx is None:
            sessions_attended_idx = 28
        if no_of_days_attended_idx is None:
            no_of_days_attended_idx = 29
        if day_wise_attendance_pct_idx is None:
            day_wise_attendance_pct_idx = 30
        if session_wise_attendance_pct_idx is None:
            session_wise_attendance_pct_idx = 31

        # Locate the optional Email column
        email_idx = None
        for idx, h_val in enumerate(header_row):
            if h_val and "email" in h_val.lower():
                email_idx = idx
                break

        # The subjects limit index in Row 7 (the boundary where grades end and summaries start)
        subjects_limit_idx = len(header_row)
        for idx, h_val in enumerate(header_row):
            if h_val and "skillfy app downloaded" in h_val.lower():
                subjects_limit_idx = idx
                break

        # Find where subject codes and Skillfy stats end
        # We can dynamically identify subjects from column 3 up to subjects_limit_idx
        subjects = []
        for idx in range(3, subjects_limit_idx):
            h_val = header_row[idx]
            # Exclude the Email column from subjects list if it exists
            if h_val and idx != email_idx:
                subjects.append(h_val)

        consecutive_empty = 0
        for r_idx in range(8, sheet.max_row + 1):
            row = sheet[r_idx]
            if not any(cell.value is not None for cell in row):
                consecutive_empty += 1
                if consecutive_empty > 50:
                    break
                continue
                
            try:
                # 1. Legend termination check to break early
                row_str_vals = [str(row[i].value).strip().upper() if row[i].value is not None else "" for i in range(min(len(row), 4))]
                is_legend = False
                for val in row_str_vals:
                    if val in [
                        "CORE EMPLOYIBILITY SKILLS", "GROW ORIENTATION", "ABSENT", 
                        "SELF-PLACEMENT", "EARLY-PLACEMENT", "QAE SKILL SHIFT VERIFICATION", 
                        "ADVANCE EXCEL", "TRAINING CLOSURE REPORT", "NAME", "SR NO.", "ASP ID"
                    ] or any(k in val for k in ["CES*", "GO*", "AB*", "SP*", "EP*", "QAE*"]):
                        is_legend = True
                        break
                if is_legend:
                    break

                name_val = row[name_idx].value
                asp_id_val = row[asp_id_idx].value
                sr_no_val = row[0].value
                email_val = row[email_idx].value if email_idx is not None else None
                
                if not name_val:
                    consecutive_empty += 1
                    if consecutive_empty > 50:
                        break
                    continue
                
                consecutive_empty = 0
                name_str = str(name_val).strip()

                # Determine if it's a student row (either Sr No or Asp ID must be numeric)
                is_numeric_sr_no = False
                if sr_no_val is not None:
                    try:
                        float(str(sr_no_val).strip())
                        is_numeric_sr_no = True
                    except ValueError:
                        pass
                
                is_numeric_asp_id = False
                if asp_id_val is not None:
                    try:
                        float(str(asp_id_val).strip())
                        is_numeric_asp_id = True
                    except ValueError:
                        pass

                if not is_numeric_sr_no and not is_numeric_asp_id:
                    continue

                name = name_str
                asp_id = None
                if asp_id_val is not None and str(asp_id_val).strip() != "":
                    try:
                        asp_id = int(float(str(asp_id_val).strip()))
                    except ValueError:
                        pass
                
                # Match student by Email first if available, else by Name
                student = None
                if email_val and str(email_val).strip() != "":
                    student_email_val = str(email_val).strip().lower()
                    student = await db.users.find_one({
                        "email": student_email_val,
                        "role": "student"
                    })
                
                if not student:
                    # Match student by Name and Role, ideally in this batch
                    student = await db.users.find_one({
                        "name": {"$regex": f"^{name}$", "$options": "i"},
                        "role": "student",
                        "batch_id": batch_id
                    })
                    if not student:
                        # Fallback to search by Name globally
                        student = await db.users.find_one({
                            "name": {"$regex": f"^{name}$", "$options": "i"},
                            "role": "student"
                        })
                
                if student:
                    student_email = student["email"].lower()
                    # Sync asp_id
                    if asp_id and student.get("asp_id") != asp_id:
                        await db.users.update_one({"_id": student["_id"]}, {"$set": {"asp_id": asp_id}})
                else:
                    # Fallback email if student is not registered yet
                    sanitized_name = "".join(c for c in name.lower() if c.isalnum())
                    student_email = f"{sanitized_name}@academy.com"
                    errors.append(f"Row {r_idx}: Student '{name}' not found in DB. Stored records under fallback '{student_email}'.")

                # Extract grades/scores
                grades = {}
                for col_idx in range(3, subjects_limit_idx):
                    if col_idx == email_idx:
                        continue
                    h_name = header_row[col_idx]
                    if h_name:
                        grades[h_name] = row[col_idx].value

                # Extract Skillfy & stats columns (using dynamically located indexes)
                skillfy_downloaded = "None"
                skillfy_courses_completed = 0
                total_sessions = 0
                sessions_attended = 0
                no_of_days_attended = 0
                day_wise_attendance_pct = 0.0
                session_wise_attendance_pct = 0.0

                if skillfy_downloaded_idx is not None and skillfy_downloaded_idx < len(row):
                    skillfy_downloaded = str(row[skillfy_downloaded_idx].value or "None").strip()
                if skillfy_courses_completed_idx is not None and skillfy_courses_completed_idx < len(row):
                    skillfy_courses_completed = safe_int(row[skillfy_courses_completed_idx].value)
                if total_sessions_idx is not None and total_sessions_idx < len(row):
                    total_sessions = safe_int(row[total_sessions_idx].value)
                if sessions_attended_idx is not None and sessions_attended_idx < len(row):
                    sessions_attended = safe_int(row[sessions_attended_idx].value)
                if no_of_days_attended_idx is not None and no_of_days_attended_idx < len(row):
                    no_of_days_attended = safe_int(row[no_of_days_attended_idx].value)
                if day_wise_attendance_pct_idx is not None and day_wise_attendance_pct_idx < len(row):
                    day_wise_attendance_pct = safe_float(row[day_wise_attendance_pct_idx].value)
                if session_wise_attendance_pct_idx is not None and session_wise_attendance_pct_idx < len(row):
                    session_wise_attendance_pct = safe_float(row[session_wise_attendance_pct_idx].value)

                # Save / Upsert Student Attendance Summary
                await db.student_attendance_summaries.update_one(
                    {
                        "student_email": student_email,
                        "batch_id": batch_id
                    },
                    {
                        "$set": {
                            "name": name,
                            "asp_id": asp_id,
                            "grades": grades,
                            "skillfy_downloaded": skillfy_downloaded,
                            "skillfy_courses_completed": skillfy_courses_completed,
                            "total_sessions": total_sessions,
                            "sessions_attended": sessions_attended,
                            "no_of_days_attended": no_of_days_attended,
                            "day_wise_attendance_pct": day_wise_attendance_pct,
                            "session_wise_attendance_pct": session_wise_attendance_pct,
                            "updated_at": datetime.now(timezone.utc),
                            "uploaded_by": current_user["email"]
                        }
                    },
                    upsert=True
                )
                summaries_added += 1

                # Parse daily session attendance starting from column index 32 onwards
                for c in range(32, len(header_row)):
                    if header_row[c] == 'S-1':
                        # Resolve date in row 5
                        d_val = row_5[c]
                        offset = 0
                        while d_val is None:
                            offset += 1
                            if c - offset < 0:
                                break
                            d_val = row_5[c - offset]
                        
                        date_str = parse_date_to_str(d_val)
                        if not date_str:
                            continue

                        if user_role != "head" and date_str != today_str:
                            # Skip past/future dates for teachers
                            continue
                                
                        # Get sessions S-1 to S-4
                        s1_val = row[c].value if c < len(row) else None
                        s2_val = row[c+1].value if c+1 < len(row) else None
                        s3_val = row[c+2].value if c+2 < len(row) else None
                        s4_val = row[c+3].value if c+3 < len(row) else None
                        
                        def map_val(v):
                            if v is None or str(v).strip() in ["", "-", "None", "none"]:
                                return "None"
                            v_str = str(v).strip()
                            v_lower = v_str.lower()
                            if v_lower in ["ab", "absent"]:
                                return "Absent"
                            if v_lower == "late":
                                return "Late"
                            if v_lower == "none":
                                return "None"
                            
                            # If it has any alphabets, it is present (return the value itself)
                            if any(c.isalpha() for c in v_str):
                                return v_str
                            return "None"
                            
                        session_1 = map_val(s1_val)
                        session_2 = map_val(s2_val)
                        session_3 = map_val(s3_val)
                        session_4 = map_val(s4_val)
                        
                        existing = await db.attendance.find_one({
                            "student_email": student_email,
                            "batch_id": batch_id,
                            "date": date_str
                        })
                        if existing:
                            session_1 = existing.get("session_1", "None") if session_1 == "None" else session_1
                            session_2 = existing.get("session_2", "None") if session_2 == "None" else session_2
                            session_3 = existing.get("session_3", "None") if session_3 == "None" else session_3
                            session_4 = existing.get("session_4", "None") if session_4 == "None" else session_4

                        await db.attendance.update_one(
                            {
                                "student_email": student_email,
                                "batch_id": batch_id,
                                "date": date_str
                            },
                            {
                                "$set": {
                                    "session_1": session_1,
                                    "session_2": session_2,
                                    "session_3": session_3,
                                    "session_4": session_4,
                                    "uploaded_by": current_user["email"]
                                }
                            },
                            upsert=True
                        )
                        records_added += 1

            except Exception as err:
                errors.append(f"Row {r_idx}: Formatting error. Details: {str(err)}")

        return {
            "message": f"Successfully processed {summaries_added} student summaries and {records_added} attendance records (new format).",
            "warnings_errors": errors
        }

@router.get("/my-attendance", response_model=StudentAttendanceSummary)
async def get_my_attendance(
    db: AsyncIOMotorDatabase = Depends(get_db),
    current_user = Depends(require_role(["student"]))
):
    """Retrieve logged-in student's daily attendance records."""
    student_email = current_user["email"].lower()
    cursor = db.attendance.find({"student_email": student_email}).sort("date", -1)
    
    records = []
    async for doc in cursor:
        doc["id"] = str(doc["_id"])
        records.append(doc)
        
    total_conducted_sessions = 0
    attended_sessions = 0
    total_conducted_days = 0
    attended_days = 0
    
    for r in records:
        day_conducted = False
        day_present = False
        
        for sess_key in ["session_1", "session_2", "session_3", "session_4"]:
            val = r.get(sess_key, "None")
            if val and val != "None":
                total_conducted_sessions += 1
                day_conducted = True
                if val != "Absent":
                    attended_sessions += 1
                    day_present = True
                    
        if day_conducted:
            total_conducted_days += 1
            if day_present:
                attended_days += 1
                
    day_wise_attendance_pct = round((attended_days / total_conducted_days) * 100, 2) if total_conducted_days > 0 else 0.0
    session_wise_attendance_pct = round((attended_sessions / total_conducted_sessions) * 100, 2) if total_conducted_sessions > 0 else 0.0
    
    return {
        "student_email": student_email,
        "records": records,
        "day_wise_attendance_pct": day_wise_attendance_pct,
        "session_wise_attendance_pct": session_wise_attendance_pct,
        "total_conducted_sessions": total_conducted_sessions,
        "attended_sessions": attended_sessions,
        "total_conducted_days": total_conducted_days,
        "attended_days": attended_days
    }

@router.get("/batch/{batch_id}", response_model=List[AttendanceRecord])
async def get_batch_attendance(
    batch_id: str,
    date: Optional[str] = None,
    db: AsyncIOMotorDatabase = Depends(get_db),
    current_user = Depends(require_role(["head", "trainer", "associate"]))
):
    """List attendance logs for a specific batch. Optionally filter by date. Trainers are restricted to assigned classes."""
    user_role = current_user.get("role")
    if user_role not in ["head", "associate"]:
        assigned_classes = current_user.get("classes_assigned") or []
        if batch_id not in assigned_classes:
            raise HTTPException(
                status_code=status.HTTP_403_FORBIDDEN,
                detail=f"You are not assigned to class/batch '{batch_id}'."
            )
            
    query = {"batch_id": batch_id}
    if date:
        query["date"] = date
        
    cursor = db.attendance.find(query).sort("date", -1)
    records = []
    async for doc in cursor:
        doc["id"] = str(doc["_id"])
        records.append(doc)
    return records

@router.get("/batch/{batch_id}/students", response_model=List[Dict[str, Any]])
async def get_batch_students(
    batch_id: str,
    db: AsyncIOMotorDatabase = Depends(get_db),
    current_user = Depends(require_role(["head", "trainer", "associate"]))
):
    """Retrieve all students assigned to a specific batch."""
    # Authenticate trainer
    user_role = current_user.get("role")
    if user_role not in ["head", "associate"]:
        assigned_classes = current_user.get("classes_assigned") or []
        if batch_id not in assigned_classes:
            raise HTTPException(
                status_code=status.HTTP_403_FORBIDDEN,
                detail=f"You are not assigned to class/batch '{batch_id}'."
            )
            
    cursor = db.users.find({"role": "student", "batch_id": batch_id}).sort("name", 1)
    students = []
    async for doc in cursor:
        students.append({
            "email": doc["email"],
            "name": doc.get("name") or doc["email"],
            "asp_id": doc.get("asp_id")
        })
    return students

@router.post("/manual-update")
async def save_manual_attendance(
    payload: AttendanceManualUpdateRequest,
    db: AsyncIOMotorDatabase = Depends(get_db),
    current_user = Depends(require_role(["head", "trainer"]))
):
    """
    Save/update attendance manually for a batch and date.
    Trainers can only update for batches assigned to them.
    """
    batch_id = payload.batch_id
    user_role = current_user.get("role")
    if user_role != "head":
        assigned_classes = current_user.get("classes_assigned") or []
        if batch_id not in assigned_classes:
            raise HTTPException(
                status_code=status.HTTP_403_FORBIDDEN,
                detail=f"You are not assigned to class/batch '{batch_id}'."
            )

    # Validate date format (YYYY-MM-DD)
    try:
        datetime.strptime(payload.date, "%Y-%m-%d")
    except ValueError:
        raise HTTPException(status_code=400, detail="Invalid date format. Expected YYYY-MM-DD.")

    # Only head admin can change past or future attendance. Trainers can only update current date attendance.
    today_str = datetime.now().strftime("%Y-%m-%d")
    if user_role != "head" and payload.date != today_str:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail=f"Only the Area Head has permission to update past or future attendance records. Current date is {today_str}."
        )

    records_updated = 0
    valid_states = ["Present", "Absent", "Late", "None"]

    for r in payload.records:
        student_email = r.student_email.lower().strip()
        
        session_1 = r.session_1 if r.session_1 in valid_states else "None"
        session_2 = r.session_2 if r.session_2 in valid_states else "None"
        session_3 = r.session_3 if r.session_3 in valid_states else "None"
        session_4 = r.session_4 if r.session_4 in valid_states else "None"

        await db.attendance.update_one(
            {
                "student_email": student_email,
                "batch_id": batch_id,
                "date": payload.date
            },
            {
                "$set": {
                    "session_1": session_1,
                    "session_2": session_2,
                    "session_3": session_3,
                    "session_4": session_4,
                    "uploaded_by": current_user["email"]
                }
            },
            upsert=True
        )
        records_updated += 1

    return {
        "message": f"Successfully updated attendance for {records_updated} students on {payload.date}."
    }

@router.get("/batch/{batch_id}/summary")
async def get_batch_attendance_summary(
    batch_id: str,
    db: AsyncIOMotorDatabase = Depends(get_db),
    current_user = Depends(require_role(["head", "trainer", "associate"]))
):
    """Retrieve overall student statistics, grades, and Skillfy app details for a batch."""
    user_role = current_user.get("role")
    if user_role not in ["head", "associate"]:
        assigned_classes = current_user.get("classes_assigned") or []
        if batch_id not in assigned_classes:
            raise HTTPException(
                status_code=status.HTTP_403_FORBIDDEN,
                detail=f"You are not assigned to class/batch '{batch_id}'."
            )
            
    cursor = db.student_attendance_summaries.find({"batch_id": batch_id}).sort("name", 1)
    summaries = []
    async for doc in cursor:
        doc["id"] = str(doc["_id"])
        summaries.append(doc)
    return summaries

@router.get("/batch/{batch_id}/stats")
async def get_batch_attendance_stats(
    batch_id: str,
    db: AsyncIOMotorDatabase = Depends(get_db),
    current_user = Depends(require_role(["head", "trainer", "associate"]))
):
    """
    Returns daywise attendance stats, topics taught per day,
    and the total count of classes conducted for each topic.
    """
    # 1. Authorize trainer check
    user_role = current_user.get("role")
    if user_role not in ["head", "associate"]:
        assigned_classes = current_user.get("classes_assigned") or []
        if batch_id not in assigned_classes:
            raise HTTPException(
                status_code=status.HTTP_403_FORBIDDEN,
                detail=f"You are not assigned to class/batch '{batch_id}'."
            )

    cursor = db.attendance.find({"batch_id": batch_id})
    records = []
    async for doc in cursor:
        records.append(doc)

    # Group by date
    date_groups = {}
    classes_conducted = set()  # set of (date, session_index, topic)

    for r in records:
        date = r["date"]
        if date not in date_groups:
            date_groups[date] = {
                "date": date,
                "presents": 0,
                "absents": 0,
                "topics": set(),
                "students_count": 0
            }

        # Check if student was present (any session is not Absent and not None)
        is_present = False
        is_absent = False
        
        for idx, sess_key in enumerate(["session_1", "session_2", "session_3", "session_4"], 1):
            val = r.get(sess_key, "None")
            if val not in ["Absent", "None"]:
                is_present = True
                if val != "Present" and val != "Late":
                    date_groups[date]["topics"].add(val)
                    classes_conducted.add((date, idx, val))
            elif val == "Absent":
                is_absent = True
        
        if is_present:
            date_groups[date]["presents"] += 1
        elif is_absent:
            date_groups[date]["absents"] += 1
            
        date_groups[date]["students_count"] += 1

    # Format day summaries
    day_summaries = []
    for date, info in sorted(date_groups.items(), key=lambda x: x[0], reverse=True):
        day_summaries.append({
            "date": date,
            "presents": info["presents"],
            "absents": info["absents"],
            "total_students": info["students_count"],
            "topics": list(info["topics"])
        })

    # Count classes per topic
    topic_counts = {}
    for date, sess_idx, topic in classes_conducted:
        topic_counts[topic] = topic_counts.get(topic, 0) + 1

    return {
        "day_summaries": day_summaries,
        "topic_counts": topic_counts
    }

@router.delete("/batch/{batch_id}/all")
async def delete_all_batch_attendance(
    batch_id: str,
    db: AsyncIOMotorDatabase = Depends(get_db),
    current_user = Depends(require_role(["head"]))
):
    """Delete all attendance records for a batch."""
    result = await db.attendance.delete_many({"batch_id": batch_id})
    
    # Reset/Delete student summaries for this batch as well
    await db.student_attendance_summaries.delete_many({"batch_id": batch_id})
    
    return {
        "message": f"Successfully deleted all {result.deleted_count} attendance records and summaries for batch '{batch_id}'."
    }

@router.delete("/batch/{batch_id}/date/{date_str}")
async def delete_batch_attendance_for_date(
    batch_id: str,
    date_str: str,
    db: AsyncIOMotorDatabase = Depends(get_db),
    current_user = Depends(require_role(["head"]))
):
    """Delete attendance records for a batch on a specific date (YYYY-MM-DD)."""
    try:
        datetime.strptime(date_str, "%Y-%m-%d")
    except ValueError:
        raise HTTPException(status_code=400, detail="Invalid date format. Expected YYYY-MM-DD.")
        
    result = await db.attendance.delete_many({
        "batch_id": batch_id,
        "date": date_str
    })
    
    return {
        "message": f"Successfully deleted {result.deleted_count} attendance records for batch '{batch_id}' on {date_str}."
    }


