import io
import json
import uuid
import logging
from datetime import datetime, timezone
from typing import List, Dict, Any, Optional
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form, status
from fastapi.responses import StreamingResponse, Response
from motor.motor_asyncio import AsyncIOMotorDatabase
from bson import ObjectId

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.core.database import get_db
from app.modules.auth.service import get_current_user
from app.modules.ai_features.schemas import (
    LLMConfigSchema,
    LLMConfigResponse,
    ATSAnalysisRequest,
    ATSAnalysisResponse,
    ResumeSaveRequest,
    ResumeSaveResponse,
    TailorResumeRequest,
    TailorResumeResponse,
    InterviewStartRequest,
    InterviewStartResponse,
    InterviewRespondRequest,
    InterviewRespondResponse,
    InterviewFeedback,
    CoachChatRequest,
    CoachChatResponse,
    RAGRecommendation,
    TestResponse,
    TestSubmitRequest,
    TestSubmitResponse,
    TestResultsResponse
)
from app.modules.ai_features.llm_client import call_llm
from app.modules.ai_features.agent_workflow import (
    parse_pdf_text,
    run_ats_analyzer_agentic_workflow,
    run_rag_search
)

router = APIRouter(prefix="/ai-features", tags=["AI Placement & Career Assistant"])
logger = logging.getLogger("ai_router")

# --- Helper functions ---

def mask_key(key: str) -> str:
    if not key:
        return ""
    if len(key) <= 8:
        return "****"
    return f"****{key[-4:]}"

async def get_user_llm_config_or_raise(db: AsyncIOMotorDatabase, email: str):
    config = await db.user_llm_configs.find_one({"user_email": email.lower()})
    if not config or not config.get("active_provider"):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="LLM settings not configured. Please go to AI Settings and enter your API keys first."
        )
    active_provider = config["active_provider"]
    api_keys = config.get("api_keys", {})
    api_key = api_keys.get(active_provider)
    if not api_key:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"API key for active provider '{active_provider}' is not configured."
        )
    return active_provider, api_key

def escape_latex(text: str) -> str:
    if not text:
        return ""
    replacements = {
        "\\": "\\textbackslash{}",
        "&": "\\&",
        "%": "\\%",
        "$": "\\$",
        "#": "\\#",
        "_": "\\_",
        "{": "\\{",
        "}": "\\}",
        "~": "\\textasciitilde{}",
        "^": "\\textasciicircum{}",
    }
    # Make sure backslash is replaced first
    text = text.replace("\\", "\\textbackslash{}")
    for char, rep in replacements.items():
        if char != "\\":
            text = text.replace(char, rep)
    return text

def generate_latex_source(data: Dict[str, Any]) -> str:
    personal = data.get("personal_info", {})
    summary = data.get("summary", "")
    experience = data.get("experience", [])
    education = data.get("education", [])
    skills = data.get("skills", [])
    projects = data.get("projects", [])
    certifications = data.get("certifications", [])
    template_name = data.get("template", "modern").lower()

    # Pre-compile parts for Custom Template placeholders
    comp_name = escape_latex(personal.get("name", ""))
    comp_title = escape_latex(personal.get("title", ""))
    comp_email = escape_latex(personal.get("email", ""))
    comp_phone = escape_latex(personal.get("phone", ""))
    comp_location = escape_latex(personal.get("location", ""))
    comp_website = escape_latex(personal.get("website", ""))
    
    comp_summary = escape_latex(summary)
    
    comp_exp = ""
    for exp in experience:
        comp_exp += r"\textbf{" + escape_latex(exp.get("role", "")) + r"} \hfill " + escape_latex(exp.get("startDate", "")) + " -- " + escape_latex(exp.get("endDate", "")) + r" \\" + "\n"
        comp_exp += r"\textit{" + escape_latex(exp.get("company", "")) + r"} \hfill " + escape_latex(exp.get("location", "")) + "\n"
        comp_exp += r"\begin{itemize}[noitemsep,topsep=2pt,leftmargin=15pt]" + "\n"
        bullets = exp.get("description", "").split("\n")
        for bullet in bullets:
            bullet = bullet.strip().lstrip("-*•").strip()
            if bullet:
                comp_exp += r"    \item " + escape_latex(bullet) + "\n"
        comp_exp += r"\end{itemize}" + "\n"
        comp_exp += r"\vspace{5pt}" + "\n"

    comp_edu = ""
    for edu in education:
        comp_edu += r"\textbf{" + escape_latex(edu.get("institution", "")) + r"} \hfill " + escape_latex(edu.get("startDate", "")) + " -- " + escape_latex(edu.get("endDate", "")) + r" \\" + "\n"
        comp_edu += r"\textit{" + escape_latex(edu.get("degree", "")) + " in " + escape_latex(edu.get("major", "")) + r"} \hfill " + escape_latex(edu.get("gpa", "")) + "\n"
        comp_edu += r"\vspace{5pt}" + "\n"

    comp_skills = r"\begin{description}[noitemsep,leftmargin=0pt]" + "\n"
    for sk in skills:
        category = sk.get("category", "General")
        skills_list = sk.get("list", "")
        comp_skills += r"    \item[\textbf{" + escape_latex(category) + r"}:] " + escape_latex(skills_list) + "\n"
    comp_skills += r"\end{description}"

    comp_proj = ""
    for proj in projects:
        title = proj.get("title", "")
        role = proj.get("role", "")
        link = proj.get("link", "")
        desc = proj.get("description", "")
        title_str = title
        if link:
            title_str = r"\href{" + link + r"}{" + title + r"}"
        comp_proj += r"\textbf{" + title_str + r"}"
        if role:
            comp_proj += r" -- \textit{" + escape_latex(role) + r"}"
        comp_proj += r" \\" + "\n" + escape_latex(desc) + r"\vspace{5pt}" + "\n"

    comp_certs = r"\begin{itemize}[noitemsep,topsep=2pt,leftmargin=15pt]" + "\n"
    for cert in certifications:
        name = cert.get("name", "")
        issuer = cert.get("issuer", "")
        date = cert.get("date", "")
        cert_str = name
        if issuer:
            cert_str += f" - {issuer}"
        if date:
            cert_str += f" ({date})"
        comp_certs += r"    \item " + escape_latex(cert_str) + "\n"
    comp_certs += r"\end{itemize}"

    # If user selected custom template and provided markup
    if template_name == "custom":
        custom_latex = data.get("custom_latex_template", "")
        if custom_latex and custom_latex.strip():
            # Replacements
            result = custom_latex
            result = result.replace("{{NAME}}", comp_name)
            result = result.replace("{{TITLE}}", comp_title)
            result = result.replace("{{EMAIL}}", comp_email)
            result = result.replace("{{PHONE}}", comp_phone)
            result = result.replace("{{LOCATION}}", comp_location)
            result = result.replace("{{WEBSITE}}", comp_website)
            result = result.replace("{{SUMMARY}}", comp_summary)
            result = result.replace("{{EXPERIENCE}}", comp_exp)
            result = result.replace("{{EDUCATION}}", comp_edu)
            result = result.replace("{{SKILLS}}", comp_skills)
            result = result.replace("{{PROJECTS}}", comp_proj)
            result = result.replace("{{CERTIFICATIONS}}", comp_certs)
            return result

    # Standard / Professional / Academic Predefined Templates
    if template_name == "professional":
        # Left-Aligned layout with Navy Accent Color
        latex = r"""\documentclass[10pt,letterpaper]{article}
\usepackage[utf8]{inputenc}
\usepackage[margin=0.75in]{geometry}
\usepackage{titlesec}
\usepackage{enumitem}
\usepackage{xcolor}
\usepackage[colorlinks=true,urlcolor=blue]{hyperref}

\definecolor{primary}{HTML}{1A365D}
\pagestyle{empty}

% Formatting section headers
\titleformat{\section}{\large\bfseries\color{primary}}{}{0em}{}[\titlerule]
\titlespacing{\section}{0pt}{10pt}{5pt}

\begin{document}

% Header Left-Aligned
\noindent
{\Huge \bfseries \color{primary} """ + comp_name + r"""} \\
{\large \textit{""" + comp_title + r"""}} \\
\vspace{4pt}
Email: \href{mailto:""" + personal.get("email", "") + r"""}{""" + comp_email + r"""} | 
Phone: """ + comp_phone + r""" | 
Location: """ + comp_location + r"""
"""
        if personal.get("website"):
            latex += r""" | Website: \href{""" + personal.get("website", "") + r"""}{""" + comp_website + r"""}"""
        latex += r"""
\vspace{10pt}
"""
        if summary:
            latex += r"""\section{Professional Summary}
""" + comp_summary + r"""
"""
        if experience:
            latex += r"""\section{Work Experience}
""" + comp_exp
        if education:
            latex += r"""\section{Education}
""" + comp_edu
        if skills:
            latex += r"""\section{Skills \& Expertises}
""" + comp_skills
        if projects:
            latex += r"""\section{Key Projects}
""" + comp_proj
        if certifications:
            latex += r"""\section{Certifications}
""" + comp_certs
        latex += r"""\end{document}"""
        return latex

    elif template_name == "academic":
        # Classic minimal serif layout
        latex = r"""\documentclass[11pt,letterpaper]{article}
\usepackage[utf8]{inputenc}
\usepackage[margin=1.0in]{geometry}
\usepackage{titlesec}
\usepackage{enumitem}
\usepackage[colorlinks=true,urlcolor=black]{hyperref}

\pagestyle{empty}

% Formatting section headers in small caps
\titleformat{\section}{\large\scshape\bfseries}{}{0em}{}[\titlerule]
\titlespacing{\section}{0pt}{12pt}{6pt}

\begin{document}

% Header
\begin{center}
    {\LARGE \scshape """ + comp_name + r"""} \\
    \vspace{2pt}
    """ + comp_title + r""" \\
    \vspace{4pt}
    Email: """ + comp_email + r""" | Phone: """ + comp_phone + r""" | Location: """ + comp_location + r"""
"""
        if personal.get("website"):
            latex += r"""    \\ Website: """ + comp_website + r"""
"""
        latex += r"""\end{center}
\vspace{8pt}
"""
        if summary:
            latex += r"""\section{Summary of Expertise}
""" + comp_summary + r"""
"""
        if experience:
            latex += r"""\section{Research \& Employment History}
""" + comp_exp
        if education:
            latex += r"""\section{Academic Qualifications}
""" + comp_edu
        if skills:
            latex += r"""\section{Technical Core Skills}
""" + comp_skills
        if projects:
            latex += r"""\section{Selected Publications \& Projects}
""" + comp_proj
        if certifications:
            latex += r"""\section{Professional Credentials}
""" + comp_certs
        latex += r"""\end{document}"""
        return latex

    elif template_name == "creative":
        # Left-Aligned with Teal Accent and customized margins
        latex = r"""\documentclass[10pt,letterpaper]{article}
\usepackage[utf8]{inputenc}
\usepackage[margin=0.75in]{geometry}
\usepackage{titlesec}
\usepackage{enumitem}
\usepackage{xcolor}
\usepackage[colorlinks=true,urlcolor=teal]{hyperref}

\definecolor{primary}{HTML}{0D9488}
\pagestyle{empty}

% Formatting section headers
\titleformat{\section}{\large\bfseries\color{primary}}{}{0em}{}[\titlerule]
\titlespacing{\section}{0pt}{10pt}{5pt}

\begin{document}

\noindent
{\Huge \bfseries \color{primary} """ + comp_name + r"""} \\
{\large \textit{""" + comp_title + r"""}} \\
\vspace{4pt}
Email: \href{mailto:""" + personal.get("email", "") + r"""}{""" + comp_email + r"""} | 
Phone: """ + comp_phone + r""" | 
Location: """ + comp_location + r"""
"""
        if personal.get("website"):
            latex += r""" | Website: \href{""" + personal.get("website", "") + r"""}{""" + comp_website + r"""}"""
        latex += r"""
\vspace{10pt}
"""
        if summary:
            latex += r"""\section{Creative Profile Summary}
""" + comp_summary + r"""
"""
        if experience:
            latex += r"""\section{Professional Experience}
""" + comp_exp
        if education:
            latex += r"""\section{Education \& Credentials}
""" + comp_edu
        if skills:
            latex += r"""\section{Core Skills \& Tools}
""" + comp_skills
        if projects:
            latex += r"""\section{Selected Projects}
""" + comp_proj
        if certifications:
            latex += r"""\section{Certifications}
""" + comp_certs
        latex += r"""\end{document}"""
        return latex

    elif template_name == "executive":
        # Centered Burgundy themed elegant serif style
        latex = r"""\documentclass[10pt,letterpaper]{article}
\usepackage[utf8]{inputenc}
\usepackage[margin=0.75in]{geometry}
\usepackage{titlesec}
\usepackage{enumitem}
\usepackage{xcolor}
\usepackage[colorlinks=true,urlcolor=brown]{hyperref}

\definecolor{primary}{HTML}{991B1B}
\pagestyle{empty}

% Formatting section headers
\titleformat{\section}{\large\bfseries\color{primary}\centerline}{}{0em}{}
\titlespacing{\section}{0pt}{12pt}{6pt}

\begin{document}

\begin{center}
    {\Huge \bfseries \color{primary} """ + comp_name + r"""} \\
    \vspace{3pt}
    {\large \scshape """ + comp_title + r"""} \\
    \vspace{4pt}
    Email: \href{mailto:""" + personal.get("email", "") + r"""}{""" + comp_email + r"""} | 
    Phone: """ + comp_phone + r""" | 
    Location: """ + comp_location + r"""
"""
        if personal.get("website"):
            latex += r""" \\ Website: \href{""" + personal.get("website", "") + r"""}{""" + comp_website + r"""}"""
        latex += r"""
\end{center}
\vspace{4pt}
"""
        if summary:
            latex += r"""
\centerline{\textbf{\color{primary} EXECUTIVE SUMMARY}}
\vspace{2pt}
\noindent
""" + comp_summary + r"""
"""
        if experience:
            latex += r"""
\vspace{8pt}
\centerline{\textbf{\color{primary} CHRONOLOGY OF EXPERIENCE}}
\vspace{2pt}
""" + comp_exp
        if education:
            latex += r"""
\vspace{8pt}
\centerline{\textbf{\color{primary} ACADEMIC FOUNDATIONS}}
\vspace{2pt}
""" + comp_edu
        if skills:
            latex += r"""
\vspace{8pt}
\centerline{\textbf{\color{primary} TECHNICAL CORE SKILLS}}
\vspace{2pt}
""" + comp_skills
        if projects:
            latex += r"""
\vspace{8pt}
\centerline{\textbf{\color{primary} SELECTED LECTURE PROJECTS}}
\vspace{2pt}
""" + comp_proj
        if certifications:
            latex += r"""
\vspace{8pt}
\centerline{\textbf{\color{primary} CERTIFICATIONS}}
\vspace{2pt}
""" + comp_certs
        latex += r"""\end{document}"""
        return latex

    else:
        # Default 'modern' centered layout
        latex = r"""\documentclass[10pt,letterpaper]{article}
\usepackage[utf8]{inputenc}
\usepackage[margin=0.75in]{geometry}
\usepackage{titlesec}
\usepackage{enumitem}
\usepackage[colorlinks=true,urlcolor=blue]{hyperref}

\pagestyle{empty}

% Formatting section headers
\titleformat{\section}{\large\bfseries}{}{0em}{}[\titlerule]
\titlespacing{\section}{0pt}{10pt}{5pt}

\begin{document}

% Header
\begin{center}
    {\Huge \bfseries """ + comp_name + r"""} \\
    \vspace{2pt}
    """ + comp_title + r""" \\
    \vspace{4pt}
    Email: \href{mailto:""" + personal.get("email", "") + r"""}{""" + comp_email + r"""} | 
    Phone: """ + comp_phone + r""" | 
    Location: """ + comp_location + r"""
    """
        if personal.get("website"):
            latex += r""" \\ \href{""" + personal.get("website", "") + r"""}{""" + comp_website + r"""}"""
        latex += r"""
\end{center}
\vspace{-10pt}
"""
        if summary:
            latex += r"""\section{Professional Summary}
""" + comp_summary + r"""
"""
        if experience:
            latex += r"""\section{Work Experience}
""" + comp_exp
        if education:
            latex += r"""\section{Education}
""" + comp_edu
        if skills:
            latex += r"""\section{Skills}
""" + comp_skills
        if projects:
            latex += r"""\section{Projects}
""" + comp_proj
        if certifications:
            latex += r"""\section{Certifications}
""" + comp_certs
        latex += r"""\end{document}"""
        return latex

def generate_word_document(data: Dict[str, Any]) -> docx.Document:
    personal = data.get("personal_info", {})
    summary = data.get("summary", "")
    experience = data.get("experience", [])
    education = data.get("education", [])
    skills = data.get("skills", [])
    projects = data.get("projects", [])
    certifications = data.get("certifications", [])

    doc = docx.Document()
    
    # Configure margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Style definitions
    # Name header
    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_name = p_name.add_run(personal.get("name", "Name"))
    run_name.bold = True
    run_name.font.size = Pt(20)
    run_name.font.name = 'Calibri'
    run_name.font.color.rgb = RGBColor(31, 78, 121)

    # Subtitle / Title
    if personal.get("title"):
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.paragraph_format.space_after = Pt(2)
        run_title = p_title.add_run(personal.get("title"))
        run_title.font.size = Pt(12)
        run_title.font.italic = True

    # Contact Details
    p_contact = doc.add_paragraph()
    p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_contact.paragraph_format.space_after = Pt(12)
    contact_parts = []
    if personal.get("email"):
        contact_parts.append(personal.get("email"))
    if personal.get("phone"):
        contact_parts.append(personal.get("phone"))
    if personal.get("location"):
        contact_parts.append(personal.get("location"))
    if personal.get("website"):
        contact_parts.append(personal.get("website"))
    
    run_contact = p_contact.add_run("  |  ".join(contact_parts))
    run_contact.font.size = Pt(9.5)
    run_contact.font.color.rgb = RGBColor(80, 80, 80)

    # Helper to add horizontal lines for headers
    def add_section_header(title_text):
        p_head = doc.add_paragraph()
        p_head.paragraph_format.space_before = Pt(12)
        p_head.paragraph_format.space_after = Pt(4)
        run_head = p_head.add_run(title_text.upper())
        run_head.bold = True
        run_head.font.size = Pt(12)
        run_head.font.color.rgb = RGBColor(31, 78, 121)
        # Add bottom border via horizontal line
        p_line = doc.add_paragraph()
        p_line.paragraph_format.space_after = Pt(6)
        p_line_run = p_line.add_run("━" * 58)
        p_line_run.font.size = Pt(6)
        p_line_run.font.color.rgb = RGBColor(200, 200, 200)

    # Summary
    if summary:
        add_section_header("Professional Summary")
        doc.add_paragraph(summary)

    # Experience
    if experience:
        add_section_header("Work Experience")
        for exp in experience:
            p_title = doc.add_paragraph()
            p_title.paragraph_format.space_after = Pt(2)
            role_run = p_title.add_run(exp.get("role", ""))
            role_run.bold = True
            
            p_title.add_run("   •   " + exp.get("company", ""))
            
            date_run = p_title.add_run(f"\t({exp.get('startDate', '')} – {exp.get('endDate', '')})")
            date_run.font.color.rgb = RGBColor(100, 100, 100)
            
            if exp.get("location"):
                p_loc = doc.add_paragraph()
                p_loc.paragraph_format.space_after = Pt(2)
                loc_run = p_loc.add_run(exp.get("location", ""))
                loc_run.font.italic = True
                loc_run.font.size = Pt(9)
                loc_run.font.color.rgb = RGBColor(120, 120, 120)
                
            bullets = exp.get("description", "").split("\n")
            for b in bullets:
                b = b.strip().lstrip("-*•").strip()
                if b:
                    doc.add_paragraph(b, style='List Bullet')

    # Education
    if education:
        add_section_header("Education")
        for edu in education:
            p_edu = doc.add_paragraph()
            p_edu.paragraph_format.space_after = Pt(2)
            inst_run = p_edu.add_run(edu.get("institution", ""))
            inst_run.bold = True
            
            degree_str = f"   •   {edu.get('degree', '')} in {edu.get('major', '')}"
            p_edu.add_run(degree_str)
            
            date_run = p_edu.add_run(f"\t({edu.get('startDate', '')} – {edu.get('endDate', '')})")
            date_run.font.color.rgb = RGBColor(100, 100, 100)
            
            if edu.get("gpa"):
                p_gpa = doc.add_paragraph()
                p_gpa.paragraph_format.space_after = Pt(4)
                gpa_run = p_gpa.add_run(f"GPA/Marks: {edu.get('gpa')}")
                gpa_run.font.size = Pt(9.5)

    # Skills
    if skills:
        add_section_header("Skills")
        for sk in skills:
            p_sk = doc.add_paragraph()
            p_sk.paragraph_format.space_after = Pt(2)
            cat_run = p_sk.add_run(sk.get("category", "General") + ": ")
            cat_run.bold = True
            p_sk.add_run(sk.get("list", ""))

    # Projects
    if projects:
        add_section_header("Projects")
        for proj in projects:
            p_proj = doc.add_paragraph()
            p_proj.paragraph_format.space_after = Pt(2)
            proj_run = p_proj.add_run(proj.get("title", ""))
            proj_run.bold = True
            
            if proj.get("role"):
                p_proj.add_run(f" ({proj.get('role')})")
            
            if proj.get("link"):
                p_proj.add_run(f"  |  Link: {proj.get('link')}")
                
            doc.add_paragraph(proj.get("description", ""))

    # Certifications
    if certifications:
        add_section_header("Certifications")
        for cert in certifications:
            name = cert.get("name", "")
            issuer = cert.get("issuer", "")
            date = cert.get("date", "")
            cert_str = name
            if issuer:
                cert_str += f" - {issuer}"
            if date:
                cert_str += f" ({date})"
            doc.add_paragraph(cert_str, style='List Bullet')

    return doc


# --- Endpoints implementation ---

@router.get("/llm-config", response_model=LLMConfigResponse)
async def get_llm_config(
    db: AsyncIOMotorDatabase = Depends(get_db),
    current_user = Depends(get_current_user)
):
    """Retrieves the list of saved API key configurations with keys masked."""
    config = await db.user_llm_configs.find_one({"user_email": current_user["email"].lower()})
    if not config:
        return LLMConfigResponse(active_provider=None, saved_providers={})
    
    saved_providers = {}
    api_keys = config.get("api_keys", {})
    for provider, key in api_keys.items():
        if key:
            saved_providers[provider] = mask_key(key)
            
    return LLMConfigResponse(
        active_provider=config.get("active_provider"),
        saved_providers=saved_providers
    )

@router.post("/llm-config", response_model=ResumeSaveResponse)
async def save_llm_config(
    config_in: LLMConfigSchema,
    db: AsyncIOMotorDatabase = Depends(get_db),
    current_user = Depends(get_current_user)
):
    """Saves or updates LLM API configurations for the user."""
    email = current_user["email"].lower()
    provider = config_in.provider.lower().strip()
    key = config_in.api_key.strip()
    
    # Validate provider name
    if provider not in ["gemini", "groq", "mistral", "openai"]:
        raise HTTPException(status_code=400, detail="Invalid provider. Must be: gemini, groq, mistral, or openai")

    existing = await db.user_llm_configs.find_one({"user_email": email})
    
    if existing:
        api_keys = existing.get("api_keys", {})
        # If the user sent a masked key and it matches our pattern, we do NOT overwrite
        if key.startswith("****"):
            if provider not in api_keys or not api_keys[provider]:
                raise HTTPException(status_code=400, detail="Cannot save masked key. Please enter the full API key.")
        else:
            api_keys[provider] = key

        await db.user_llm_configs.update_one(
            {"user_email": email},
            {"$set": {"active_provider": provider, "api_keys": api_keys, "updated_at": datetime.now(timezone.utc)}}
        )
    else:
        if key.startswith("****"):
            raise HTTPException(status_code=400, detail="Cannot save masked key. Please enter the full API key.")
        
        await db.user_llm_configs.insert_one({
            "user_email": email,
            "active_provider": provider,
            "api_keys": {provider: key},
            "updated_at": datetime.now(timezone.utc)
        })
        
    return ResumeSaveResponse(status="success", message=f"Successfully set active provider to {provider}.")


@router.post("/ats-analyze", response_model=ATSAnalysisResponse)
async def analyze_ats(
    job_description: str = Form(...),
    student_email: Optional[str] = Form(None),
    resume_file: Optional[UploadFile] = File(None),
    db: AsyncIOMotorDatabase = Depends(get_db),
    current_user = Depends(get_current_user)
):
    """
    Extracts text from the resume and runs the agentic pipeline to compare it to the JD, 
    returning ATS score, Explanations (XAI) and internal course/book recommendations (RAG).
    """
    user_email = current_user["email"].lower()
    
    # 1. Retrieve the student's resume text
    resume_text = ""
    
    if resume_file:
        # File uploaded directly
        if not resume_file.filename.lower().endswith(".pdf"):
            raise HTTPException(status_code=400, detail="Uploaded file must be a PDF.")
        file_bytes = await resume_file.read()
        resume_text = await parse_pdf_text(file_bytes)
    else:
        # Pull from DB
        target_email = student_email.lower().strip() if student_email else user_email
        
        # If trainer or admin, they can query others. If student, they can only query their own
        if current_user.get("role") == "student" and target_email != user_email:
            raise HTTPException(status_code=403, detail="Students can only analyze their own resume.")
            
        resume_doc = await db.resumes.find_one({"student_email": target_email})
        if not resume_doc:
            # Let's check if they have builder data saved
            builder_data = await db.user_resumes.find_one({"user_email": target_email})
            if builder_data:
                # Compile builder details to text for parsing
                text_parts = []
                p = builder_data.get("personal_info", {})
                text_parts.append(f"{p.get('name')} {p.get('title')}")
                text_parts.append(builder_data.get("summary", ""))
                for sk in builder_data.get("skills", []):
                    text_parts.append(f"{sk.get('category')}: {sk.get('list')}")
                for exp in builder_data.get("experience", []):
                    text_parts.append(f"{exp.get('role')} at {exp.get('company')}. {exp.get('description')}")
                for proj in builder_data.get("projects", []):
                    text_parts.append(f"{proj.get('title')}: {proj.get('description')}")
                resume_text = "\n".join(text_parts)
            else:
                raise HTTPException(
                    status_code=404, 
                    detail=f"No active resume PDF or resume builder data found for email: {target_email}."
                )
        else:
            filepath = resume_doc["filepath"]
            if filepath.startswith("http://") or filepath.startswith("https://"):
                async with httpx.AsyncClient() as client:
                    res = await client.get(filepath)
                    if res.status_code != 200:
                        raise HTTPException(status_code=400, detail="Failed to fetch resume file from Cloudinary.")
                    resume_text = await parse_pdf_text(res.content)
            else:
                if not os.path.exists(filepath):
                    raise HTTPException(status_code=404, detail="Local resume file not found.")
                with open(filepath, "rb") as f:
                    resume_text = await parse_pdf_text(f.read())
                    
    if not resume_text.strip():
        raise HTTPException(status_code=400, detail="Resume content is empty or unreadable.")

    # 2. Get LLM Credentials
    # Since trainers/admins use the tool on behalf of students, they use their OWN API keys
    active_provider, api_key = await get_user_llm_config_or_raise(db, user_email)

    # 3. Run Agentic Pipeline (Parser/Auditor -> RAG -> Explainer)
    result = await run_ats_analyzer_agentic_workflow(
        resume_text=resume_text,
        job_description=job_description,
        db=db,
        provider=active_provider,
        api_key=api_key
    )
    
    return result


@router.post("/resume-builder/save", response_model=ResumeSaveResponse)
async def save_resume_data(
    resume_in: ResumeSaveRequest,
    db: AsyncIOMotorDatabase = Depends(get_db),
    current_user = Depends(get_current_user)
):
    """Saves or updates the student's Master Resume profile details in MongoDB."""
    email = current_user["email"].lower()
    resume_dict = resume_in.model_dump()
    resume_dict["user_email"] = email
    resume_dict["updated_at"] = datetime.now(timezone.utc)
    
    await db.user_resumes.update_one(
        {"user_email": email},
        {"$set": resume_dict},
        upsert=True
    )
    return ResumeSaveResponse(status="success", message="Master resume details saved successfully.")

@router.get("/resume-builder/load", response_model=Optional[ResumeSaveRequest])
async def load_resume_data(
    student_email: Optional[str] = None,
    db: AsyncIOMotorDatabase = Depends(get_db),
    current_user = Depends(get_current_user)
):
    """Loads the Master Resume profile details for editing."""
    email = current_user["email"].lower()
    
    # Staff can load any student's data. Students can only load their own
    target_email = student_email.lower().strip() if student_email else email
    if current_user.get("role") == "student" and target_email != email:
        raise HTTPException(status_code=403, detail="Students can only load their own resume profile.")
        
    data = await db.user_resumes.find_one({"user_email": target_email})
    if not data:
        # Prepopulate email and name from auth account
        user_info = await db.users.find_one({"email": target_email})
        if not user_info:
            return None
        return ResumeSaveRequest(
            personal_info={"name": user_info.get("name", ""), "email": target_email, "phone": "", "location": "", "website": "", "title": ""},
            summary="",
            experience=[],
            education=[],
            skills=[],
            projects=[],
            certifications=[],
            template="modern"
        )
    return data


@router.post("/resume-builder/tailor", response_model=TailorResumeResponse)
async def tailor_resume(
    req: TailorResumeRequest,
    db: AsyncIOMotorDatabase = Depends(get_db),
    current_user = Depends(get_current_user)
):
    """
    Optimizer Agent: Takes the saved master resume details and customizes the professional summary, 
    experience bullet points, and project phrasing to optimize matching against the target JD.
    """
    email = current_user["email"].lower()
    master = await db.user_resumes.find_one({"user_email": email})
    if not master:
        raise HTTPException(
            status_code=400, 
            detail="No Master Resume details saved yet. Please fill in your resume details first."
        )
        
    active_provider, api_key = await get_user_llm_config_or_raise(db, email)
    
    # Strip database keys for prompt
    master_clean = {
        "personal_info": master.get("personal_info", {}),
        "summary": master.get("summary", ""),
        "experience": master.get("experience", []),
        "education": master.get("education", []),
        "skills": master.get("skills", []),
        "projects": master.get("projects", []),
        "certifications": master.get("certifications", []),
        "template": master.get("template", "modern")
    }

    tailor_instruction = (
        "You are an expert ATS Resume Optimizer Agent. Your task is to rewrite the candidate's professional "
        "summary, experience bullet points, and projects descriptions to align perfectly with the target "
        "Job Description. Highlight relevant keywords and use strong quantitative, metric-oriented phrasings. "
        "Do not edit their name, contact details, or institutional credentials (dates, degrees, companies). "
        "Do not invent new roles or fake titles. "
        "You must output a valid JSON object matching the input structure exactly, containing the tailored "
        "fields, plus an extra string field 'explanation' detailing exactly what optimizations were made and why. "
        "JSON output format:\n"
        "{\n"
        '  "tailored_resume_data": { "personal_info": ..., "summary": "rewritten...", "experience": [...], ... },\n'
        '  "explanation": "Summarized modifications details..."\n'
        "}\n"
        "Output ONLY valid JSON."
    )
    
    prompt = f"MASTER RESUME JSON:\n{json.dumps(master_clean, indent=2)}\n\nJOB DESCRIPTION:\n{req.job_description}"
    
    response = await call_llm(
        provider=active_provider,
        api_key=api_key,
        prompt=prompt,
        system_instruction=tailor_instruction,
        response_format_json=True
    )
    
    try:
        cleaned_json = response.strip()
        if cleaned_json.startswith("```"):
            lines = cleaned_json.split("\n")
            if lines[0].startswith("```json") or lines[0].startswith("```"):
                cleaned_json = "\n".join(lines[1:-1])
                
        result = json.loads(cleaned_json)
        return TailorResumeResponse(
            tailored_resume_data=result.get("tailored_resume_data", master_clean),
            explanation=result.get("explanation", "Resume optimized to match targets.")
        )
    except Exception as e:
        logger.error(f"Failed to parse tailor agent response: {response}. Error: {e}")
        return TailorResumeResponse(
            tailored_resume_data=master_clean,
            explanation=f"Fallback template returned. AI optimization encountered an parsing error: {str(e)}"
        )


@router.post("/resume-builder/export-latex")
async def export_latex(
    resume_in: ResumeSaveRequest
):
    """Compiles the resume details into LaTeX source code and returns it as a download."""
    latex_code = generate_latex_source(resume_in.model_dump())
    
    # Return as raw text stream
    return Response(
        content=latex_code,
        media_type="text/plain",
        headers={
            "Content-Disposition": "attachment; filename=resume.tex"
        }
    )

@router.post("/resume-builder/export-docx")
async def export_docx(
    resume_in: ResumeSaveRequest
):
    """Compiles the resume details into a styled Word document (.docx) using python-docx and returns it."""
    doc = generate_word_document(resume_in.model_dump())
    
    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    
    return StreamingResponse(
        file_stream,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={
            "Content-Disposition": "attachment; filename=resume.docx"
        }
    )


# --- AI Mock Interview Endpoints ---

@router.post("/interview/start", response_model=InterviewStartResponse)
async def start_mock_interview(
    req: InterviewStartRequest,
    db: AsyncIOMotorDatabase = Depends(get_db),
    current_user = Depends(get_current_user)
):
    """Starts a stateful AI mock interview session using RAG questions generated from candidate resume details and target JD."""
    email = current_user["email"].lower()
    active_provider, api_key = await get_user_llm_config_or_raise(db, email)

    # Compile user details as reference
    resume_text = "No resume profile saved."
    resume_doc = await db.user_resumes.find_one({"user_email": email})
    if resume_doc:
        resume_text = f"Skills: {resume_doc.get('skills', [])}. Experience: {resume_doc.get('experience', [])}. Projects: {resume_doc.get('projects', [])}."

    # Pre-generate RAG-based questions matching candidate's resume directly to target JD
    num_q = max(3, min(15, req.num_questions or 5))

    rag_questions_instruction = (
        f"You are a professional technical recruiter hiring for target role: {req.target_role}.\n"
        f"JOB DESCRIPTION:\n{req.job_description or 'General technical role.'}\n\n"
        f"CANDIDATE RESUME DETAILS:\n{resume_text}\n\n"
        f"Your task is to generate exactly {num_q} interview questions. These questions MUST directly match and compare the "
        "candidate's specific stated experiences, skills, and projects against the requirements of the job description "
        "(e.g., asking how they applied a skill they used in a project to solve a specific JD requirement). "
        f"You must output a valid JSON array containing exactly {num_q} question strings. Output structure:\n"
        "[\n"
        '  "Question 1...",\n'
        '  "Question 2...",\n'
        '  ...\n'
        "]\n"
        "Output ONLY valid JSON. No extra text or markdown format."
    )
    
    rag_response = await call_llm(
        provider=active_provider,
        api_key=api_key,
        prompt=f"Generate {num_q} RAG interview questions now.",
        system_instruction=rag_questions_instruction,
        response_format_json=True
    )
    
    # Parse questions list
    try:
        cleaned_rag = rag_response.strip()
        if cleaned_rag.startswith("```"):
            lines = cleaned_rag.split("\n")
            if lines[0].startswith("```json") or lines[0].startswith("```"):
                cleaned_rag = "\n".join(lines[1:-1])
        questions = json.loads(cleaned_rag)
        if not isinstance(questions, list) or len(questions) < num_q:
            raise ValueError(f"Generated question count is less than requested: {num_q}")
    except Exception as e:
        logger.warning(f"Failed to generate custom RAG questions count: {rag_response}. Error: {e}")
        # Fallback list matching the exact requested num_q count
        fallbacks = [
            f"Could you introduce yourself and tell me how your experience fits the {req.target_role} role?",
            f"What technical project in your profile are you most proud of, and how does it relate to the JD requirements?",
            f"How do you handle state management or equivalent architecture in your stack?",
            f"Can you explain how you would resolve a scalability or caching challenge in this architecture?",
            f"Do you have any questions for us regarding the technical environment described in the JD?",
            f"Can you discuss a time you had to debug a complex runtime issue under tight timelines?",
            f"How do you stay updated with changes and best practices in the {req.target_role} ecosystem?",
            f"Explain how you would write unit tests for the core logic described in this job profile."
        ]
        questions = [fallbacks[i % len(fallbacks)] for i in range(num_q)]
        for i in range(len(questions)):
            if i >= len(fallbacks):
                questions[i] = f"Question {i+1}: {questions[i]}"

    session_id = str(uuid.uuid4())
    first_question = questions[0]
    
    # Save session state
    session_doc = {
        "_id": session_id,
        "user_email": email,
        "target_role": req.target_role,
        "job_description": req.job_description,
        "rag_questions": questions,
        "chat_history": [
            {"role": "interviewer", "content": first_question}
        ],
        "created_at": datetime.now(timezone.utc),
        "is_complete": False
    }
    await db.interview_sessions.insert_one(session_doc)
    
    return InterviewStartResponse(session_id=session_id, first_question=first_question)

@router.post("/interview/respond", response_model=InterviewRespondResponse)
async def respond_mock_interview(
    req: InterviewRespondRequest,
    db: AsyncIOMotorDatabase = Depends(get_db),
    current_user = Depends(get_current_user)
):
    """Accepts candidate's answer, and queries LLM to transition to the next pre-generated RAG question or finalize feedback."""
    email = current_user["email"].lower()
    session = await db.interview_sessions.find_one({"_id": req.session_id})
    if not session:
        raise HTTPException(status_code=404, detail="Interview session not found.")
        
    if session.get("is_complete", False):
        return InterviewRespondResponse(is_complete=True, feedback=session.get("feedback"))

    active_provider, api_key = await get_user_llm_config_or_raise(db, email)
    
    chat_history = session.get("chat_history", [])
    chat_history.append({"role": "candidate", "content": req.user_response})
    
    # Calculate round details
    question_rounds = sum(1 for msg in chat_history if msg["role"] == "interviewer")
    rag_questions = session.get("rag_questions", [])
    
    if question_rounds >= len(rag_questions):
        # Complete interview and compile feedback
        system_instruction = (
            f"You are a Senior Technical Recruiter. Conducted an interview for: {session['target_role']}. "
            "Examine the chat history. Provide a complete evaluation of the candidate. "
            "Evaluate: Rating (Strong Hire, Hire, Weak Hire, No Hire), Score (0-100), overall strengths, "
            "weaknesses, written summary, and individual breakdowns of each answer. "
            "You must output a valid JSON object matching this structure EXACTLY:\n"
            "{\n"
            '  "rating": "Hire",\n'
            '  "score": 80,\n'
            '  "strengths": ["...", "..."],\n'
            '  "weaknesses": ["...", "..."],\n'
            '  "overall_summary": "Detailed feedback text...",\n'
            '  "question_breakdowns": [\n'
            '    {"question": "Q1...", "answer": "A1...", "score": 80, "critique": "..."},\n'
            '    ...\n'
            '  ]\n'
            "}\n"
            "Output ONLY valid JSON."
        )
        
        transcript = ""
        for msg in chat_history:
            role = "Interviewer" if msg["role"] == "interviewer" else "Candidate"
            transcript += f"{role}: {msg['content']}\n"
            
        feedback_response = await call_llm(
            provider=active_provider,
            api_key=api_key,
            prompt=transcript,
            system_instruction=system_instruction,
            response_format_json=True
        )
        
        try:
            cleaned_feedback = feedback_response.strip()
            if cleaned_feedback.startswith("```"):
                lines = cleaned_feedback.split("\n")
                if lines[0].startswith("```json") or lines[0].startswith("```"):
                    cleaned_feedback = "\n".join(lines[1:-1])
                    
            fb_dict = json.loads(cleaned_feedback)
            feedback = InterviewFeedback(**fb_dict)
        except Exception as e:
            logger.error(f"Failed to parse interview feedback: {feedback_response}. Error: {e}")
            feedback = InterviewFeedback(
                rating="Hire",
                score=70,
                strengths=["Completed the session"],
                weaknesses=["Parsing issues in feedback"],
                overall_summary="Completed mock interview session.",
                question_breakdowns=[]
            )
            
        await db.interview_sessions.update_one(
            {"_id": req.session_id},
            {"$set": {"chat_history": chat_history, "is_complete": True, "feedback": feedback.model_dump()}}
        )
        return InterviewRespondResponse(is_complete=True, feedback=feedback)
    else:
        # Continue asking questions from our RAG-pregenerated list
        next_target_question = rag_questions[question_rounds]
        
        system_instruction = (
            f"You are a professional technical interviewer conducting an interview for: {session['target_role']}. "
            "Acknowledge the candidate's last answer briefly, and transition directly to the next target question. "
            f"TARGET QUESTION TO POSE: {next_target_question}\n"
            "Present this question clearly. Keep it concise. Ask exactly one question. Do not skip the target question."
        )
        
        transcript = ""
        for msg in chat_history:
            role = "Interviewer" if msg["role"] == "interviewer" else "Candidate"
            transcript += f"{role}: {msg['content']}\n"
            
        next_question = await call_llm(
            provider=active_provider,
            api_key=api_key,
            prompt=transcript,
            system_instruction=system_instruction
        )
        
        chat_history.append({"role": "interviewer", "content": next_question})
        
        await db.interview_sessions.update_one(
            {"_id": req.session_id},
            {"$set": {"chat_history": chat_history}}
        )
        return InterviewRespondResponse(next_question=next_question, is_complete=False)

@router.get("/interview/history", response_model=List[Dict[str, Any]])
async def get_interview_history(
    db: AsyncIOMotorDatabase = Depends(get_db),
    current_user = Depends(get_current_user)
):
    """Lists previous mock interviews for the student."""
    email = current_user["email"].lower()
    cursor = db.interview_sessions.find({"user_email": email}).sort("created_at", -1)
    sessions = []
    async for doc in cursor:
        doc["session_id"] = doc["_id"]
        sessions.append(doc)
    return sessions


# --- AI Career Coach Chatbot Endpoints ---

@router.post("/coach/chat", response_model=CoachChatResponse)
async def chat_with_career_coach(
    req: CoachChatRequest,
    db: AsyncIOMotorDatabase = Depends(get_db),
    current_user = Depends(get_current_user)
):
    """
    RAG-based Career Coach: Analyzes student message, searches internal databases for matching 
    classes, notes, and library books, and provides context-aware guidance.
    """
    email = current_user["email"].lower()
    active_provider, api_key = await get_user_llm_config_or_raise(db, email)

    # 1. Search RAG database for matching materials using query words
    user_words = [w.strip("?,.!") for w in req.message.split() if len(w) > 3]
    rag_context = await run_rag_search(db, user_words)

    # 2. Look up student's active resume details for context personalization
    resume_context = "No resume profile saved yet."
    resume_doc = await db.user_resumes.find_one({"user_email": email})
    if resume_doc:
        resume_context = (
            f"Summary: {resume_doc.get('summary', '')}\n"
            f"Skills: {resume_doc.get('skills', [])}\n"
            f"Experience: {resume_doc.get('experience', [])}\n"
            f"Projects: {resume_doc.get('projects', [])}\n"
            f"Certifications: {resume_doc.get('certifications', [])}"
        )

    # 3. Format local resources summary
    resources_summary = ""
    if rag_context:
        resources_summary = "We found these local resources in the academy database matching the student's interests:\n"
        for idx, rec in enumerate(rag_context, 1):
            resources_summary += f"{idx}. [{rec.type.upper()}] Title: {rec.title}. Description: {rec.description}. Info: {rec.link_or_details}\n"
    
    # 4. Assemble chat prompt with history
    system_instruction = (
        "You are an AI Career Coach at the Academy. You guide students on choosing technical tracks, "
        "building skills, borrowing library books, and enrolling in class batches. "
        "Be warm, encouraging, smart, and professional. "
        "You should reference the student's active resume details to personalize your suggestions, "
        "helping them map what they already know or have done to new courses, reference books, or career tracks! "
        "Ensure you reference the local Academy resources provided (batches, notes, library books) to answer queries. "
        "If a specific skill is discussed, recommend relevant resources available in-house. "
        f"STUDENT CONTEXT: Name: {current_user.get('name', 'Student')}, Email: {email}, Role: {current_user.get('role')}.\n"
        f"STUDENT ACTIVE RESUME DETAILS:\n{resume_context}\n\n"
        f"LOCAL ACADEMY RESOURCES CONTEXT:\n{resources_summary or 'No specific matching resources found in database.'}"
    )

    # Build conversation context
    messages_payload = ""
    for msg in req.history[-6:]:  # Limit history to last 6 messages
        role = "Student" if msg.get("role") == "user" else "Coach"
        messages_payload += f"{role}: {msg.get('content')}\n"
        
    messages_payload += f"Student: {req.message}\n"
    
    coach_reply = await call_llm(
        provider=active_provider,
        api_key=api_key,
        prompt=messages_payload,
        system_instruction=system_instruction
    )
    
    return CoachChatResponse(response=coach_reply, rag_context=rag_context)


# --- AI Weekly Test Suite Endpoints ---

@router.post("/tests/create", response_model=ResumeSaveResponse)
async def create_weekly_test(
    title: str = Form(...),
    topic: str = Form(...),
    num_questions: int = Form(5),
    marks_per_question: int = Form(1),
    study_file: Optional[UploadFile] = File(None),
    db: AsyncIOMotorDatabase = Depends(get_db),
    current_user = Depends(get_current_user)
):
    """
    Trainer/Admin endpoint: Generates a test based on typed topic or uploaded study guide PDF.
    Uses AI RAG to formulate specific MCQ questions.
    """
    # Verify authority
    if current_user.get("role") not in ["head", "trainer"]:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Only trainers and head administrators can create tests."
        )
        
    user_email = current_user["email"].lower()
    active_provider, api_key = await get_user_llm_config_or_raise(db, user_email)

    source_text = f"Topic to test: {topic}"
    if study_file:
        if not study_file.filename.lower().endswith(".pdf"):
            raise HTTPException(status_code=400, detail="Uploaded study material must be a PDF file.")
        file_bytes = await study_file.read()
        import os
        # Need to import parse_pdf_text dynamically or use our agent_workflow parser
        resume_text = await parse_pdf_text(file_bytes)
        source_text = resume_text

    # Prompt LLM to create tests questions structure
    generator_instruction = (
        "You are an expert Academic Teacher Assistant. Your task is to generate a high-quality "
        "multiple-choice question (MCQ) weekly test based on the study text/topic provided. "
        f"You must create exactly {num_questions} questions. "
        "Each question must contain: question text, exactly 4 option choices (strings), the 0-indexed correct option "
        "index (0 to 3), and a clear explanation of why the correct option is right. "
        "You must output a valid JSON array matching this format EXACTLY:\n"
        "[\n"
        "  {\n"
        '    "question_text": "Question content here?",\n'
        '    "options": ["Option A", "Option B", "Option C", "Option D"],\n'
        '    "correct_option_index": 1,\n'
        '    "explanation": "Detail reason for correct answer."\n'
        "  },\n"
        "  ...\n"
        "]\n"
        "Output ONLY the JSON array. Do not include markdown code ticks or chat logs."
    )
    
    ai_response = await call_llm(
        provider=active_provider,
        api_key=api_key,
        prompt=source_text[:12000],  # Limit context buffer
        system_instruction=generator_instruction,
        response_format_json=True
    )
    
    try:
        cleaned_json = ai_response.strip()
        if cleaned_json.startswith("```"):
            lines = cleaned_json.split("\n")
            if lines[0].startswith("```json") or lines[0].startswith("```"):
                cleaned_json = "\n".join(lines[1:-1])
        questions = json.loads(cleaned_json)
        if not isinstance(questions, list) or len(questions) == 0:
            raise ValueError("AI response is not a valid list of questions.")
    except Exception as e:
        logger.error(f"Failed to parse test generator response: {ai_response}. Error: {e}")
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"AI failed to compile valid test questions. Error: {str(e)}"
        )

    test_id = str(uuid.uuid4())
    test_doc = {
        "_id": test_id,
        "title": title,
        "topic": topic,
        "num_questions": len(questions),
        "marks_per_question": marks_per_question,
        "questions": questions,
        "created_by": user_email,
        "created_at": datetime.now(timezone.utc)
    }
    await db.tests.insert_one(test_doc)
    
    return ResumeSaveResponse(status="success", message=f"Test '{title}' with {len(questions)} questions created successfully.")

@router.get("/tests")
async def list_active_tests(
    db: AsyncIOMotorDatabase = Depends(get_db),
    current_user = Depends(get_current_user)
):
    """Lists active tests. For students, hides answers and explanations."""
    cursor = db.tests.find().sort("created_at", -1)
    tests = []
    
    is_student = current_user.get("role") == "student"
    
    async for doc in cursor:
        doc["id"] = doc["_id"]
        # Format questions schema
        raw_qs = doc.get("questions", [])
        formatted_qs = []
        for i, q in enumerate(raw_qs):
            q_res = {
                "question_index": i,
                "question_text": q.get("question_text"),
                "options": q.get("options", [])
            }
            if not is_student:
                # Add private metadata for trainers
                q_res["correct_option_index"] = q.get("correct_option_index")
                q_res["explanation"] = q.get("explanation")
            formatted_qs.append(q_res)
            
        doc["questions"] = formatted_qs
        tests.append(doc)
        
    return tests

@router.get("/tests/{test_id}")
async def get_test_details(
    test_id: str,
    db: AsyncIOMotorDatabase = Depends(get_db),
    current_user = Depends(get_current_user)
):
    """Retrieves specific test details. For students, hides answers and explanations."""
    doc = await db.tests.find_one({"_id": test_id})
    if not doc:
        raise HTTPException(status_code=404, detail="Test not found.")
        
    doc["id"] = doc["_id"]
    is_student = current_user.get("role") == "student"
    
    raw_qs = doc.get("questions", [])
    formatted_qs = []
    for i, q in enumerate(raw_qs):
        q_res = {
            "question_index": i,
            "question_text": q.get("question_text"),
            "options": q.get("options", [])
        }
        if not is_student:
            q_res["correct_option_index"] = q.get("correct_option_index")
            q_res["explanation"] = q.get("explanation")
        formatted_qs.append(q_res)
        
    doc["questions"] = formatted_qs
    return doc

@router.post("/tests/{test_id}/submit", response_model=TestSubmitResponse)
async def submit_test_answers(
    test_id: str,
    submission: TestSubmitRequest,
    db: AsyncIOMotorDatabase = Depends(get_db),
    current_user = Depends(get_current_user)
):
    """Scores a student's responses, logs the score card, and returns explanations."""
    test = await db.tests.find_one({"_id": test_id})
    if not test:
        raise HTTPException(status_code=404, detail="Test not found.")
        
    student_email = current_user["email"].lower()
    student_name = current_user.get("name", "Student")
    
    questions = test.get("questions", [])
    answers = submission.answers
    
    correct_count = 0
    wrong_count = 0
    graded_results = []
    
    for idx, q in enumerate(questions):
        idx_str = str(idx)
        chosen_opt = answers.get(idx_str) # student chosen option index
        correct_opt = q.get("correct_option_index", 0)
        
        is_correct = chosen_opt == correct_opt
        if is_correct:
            correct_count += 1
        else:
            wrong_count += 1
            
        graded_results.append({
            "question_index": idx,
            "question_text": q.get("question_text"),
            "options": q.get("options", []),
            "chosen_option": chosen_opt,
            "correct_option": correct_opt,
            "is_correct": is_correct,
            "explanation": q.get("explanation", "")
        })
        
    marks = test.get("marks_per_question", 1)
    score = correct_count * marks
    max_score = len(questions) * marks
    
    # Save submission
    sub_doc = {
        "_id": str(uuid.uuid4()),
        "test_id": test_id,
        "student_email": student_email,
        "student_name": student_name,
        "score": float(score),
        "max_score": float(max_score),
        "answers": answers,
        "graded_at": datetime.now(timezone.utc)
    }
    await db.test_submissions.insert_one(sub_doc)
    
    return TestSubmitResponse(
        score=score,
        max_score=max_score,
        correct_count=correct_count,
        wrong_count=wrong_count,
        results=graded_results
    )

@router.get("/tests/{test_id}/results", response_model=TestResultsResponse)
async def get_test_submissions(
    test_id: str,
    db: AsyncIOMotorDatabase = Depends(get_db),
    current_user = Depends(get_current_user)
):
    """Trainer/Admin endpoint: Retrieves all grade listings for a specific test."""
    if current_user.get("role") not in ["head", "trainer"]:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Only trainers and administrators can view test results ledgers."
        )
        
    test = await db.tests.find_one({"_id": test_id})
    if not test:
        raise HTTPException(status_code=404, detail="Test not found.")
        
    cursor = db.test_submissions.find({"test_id": test_id}).sort("graded_at", -1)
    submissions = []
    
    async for doc in cursor:
        submissions.append({
            "id": doc["_id"],
            "student_email": doc["student_email"],
            "student_name": doc["student_name"],
            "score": doc["score"],
            "max_score": doc["max_score"],
            "graded_at": doc["graded_at"]
        })
        
    return TestResultsResponse(
        test_id=test_id,
        title=test.get("title", "Test"),
        submissions=submissions
    )
